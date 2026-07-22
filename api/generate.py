"""Vercel serverless endpoint for generating CRM/PowerBI workbooks."""

from __future__ import annotations

import cgi
import ast
import csv
from datetime import datetime, timezone
import json
import mimetypes
import os
import re
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from email.message import Message
from http.server import BaseHTTPRequestHandler
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import cr_maker  # noqa: E402
import lead_splitter  # noqa: E402
import program_a_report  # noqa: E402
import program_b_country_report  # noqa: E402
from api.file_uploads import (  # noqa: E402
    CHUNK_UPLOAD_MAX_BYTES,
    cleanup_upload,
    read_upload_metadata,
    resolve_uploaded_file,
    save_chunk,
)


APP_REPORT = "report"
APP_LEAD_SPLITTER = "lead_splitter"
APP_CR = "cr"
APP_DATABASE_CHECK = "database_check"
PROGRAM_A = "program_a"
PROGRAM_B = "program_b"
PROGRAM_C = "program_c"
PROGRAM_A_OUTPUT_FILENAME = "crm_powerbi_output.xlsx"
PROGRAM_B_OUTPUT_FILENAME = "crm_country_report.xlsx"
PROGRAM_C_OUTPUT_FILENAME = "crm_monthly_comments_output.xlsx"
REPORT_ACTION_UPLOAD_MONTHLY_CHUNK = "upload_monthly_chunk"
MAX_UPLOAD_BYTES = 45 * 1024 * 1024
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATABASE_CHECK_WEBHOOK_URL = os.environ.get(
    "DATABASE_CHECK_WEBHOOK_URL",
    "https://ammartd20.app.n8n.cloud/webhook-test/Database-check",
).strip()
DATABASE_CHECK_CALLBACK_PATH = (
    os.environ.get("DATABASE_CHECK_CALLBACK_PATH", "/api/n8n-callback").strip() or "/api/n8n-callback"
)
DATABASE_CHECK_CALLBACK_URL = os.environ.get("DATABASE_CHECK_CALLBACK_URL", "").strip()
DATABASE_CHECK_TIMEOUT_SECONDS_RAW = os.environ.get(
    "DATABASE_CHECK_TIMEOUT_SECONDS",
    "0",
).strip()
DATABASE_CHECK_PASSWORD = os.environ.get("DATABASE_CHECK_PASSWORD", "checker123456")
DATABASE_CHECK_LOG_SPREADSHEET_ID = os.environ.get(
    "DATABASE_CHECK_LOG_SPREADSHEET_ID",
    "1ng3xTwQ-SzDmD-XOfOczoGTfnwiXSUGq5CnZoFySkes",
).strip()
DATABASE_CHECK_LOG_SHEET_NAME = os.environ.get("DATABASE_CHECK_LOG_SHEET_NAME", "").strip()
DATABASE_CHECK_LOG_SERVICE_ACCOUNT_EMAIL = os.environ.get(
    "DATABASE_CHECK_LOG_SERVICE_ACCOUNT_EMAIL",
    "matservice@mitservice.iam.gserviceaccount.com",
).strip()
DATABASE_CHECK_INPUT_COLUMNS = [
    "Brand",
    "Account No",
    "Client Name",
    "Customer Status",
    "Country",
    "Current Assigned Agent",
    "Last 10 Comments",
]
DATABASE_CHECK_OUTPUT_COLUMNS = [
    *DATABASE_CHECK_INPUT_COLUMNS,
    "Suggested status",
    "Reason",
]
DATABASE_CHECK_WEBHOOK_COLUMNS = [
    "Brand",
    "Account No",
    "Last 10 Comments",
    "Customer Status",
    "Country",
    "Current Assigned Agent",
]
DATABASE_CHECK_WEBHOOK_KEY_BY_COLUMN = {
    "Brand": "brand",
    "Account No": "account no",
    "Last 10 Comments": "last 10 comments",
    "Customer Status": "customer status",
    "Country": "country",
    "Current Assigned Agent": "Agent",
}
DATABASE_CHECK_NON_ACTION_COMMENTS = {
    "NA",
    "no answer 1-5",
    "NA VM",
    "navm",
    "VM",
    "DVM",
    "PU HU",
    "puhu",
    "REJ",
    "rej call",
    "no answer",
    "no answer vm",
    "no pickup",
    "not picking",
    "did not pick",
    "hung up",
    "hang up",
    "hu",
    "cut call",
    "voicemail",
    "voice mail",
    "no reply",
    "no response",
    "no ring",
    "ringing",
    "ringing na",
    "ringing/na",
    "ringing / na !",
    "ringing but not answering",
    "ringing but rejecting",
    "beeping",
    "beep",
    "beeps",
    "na beeps",
    "no rout",
    "could not hear",
    "wrong ringing",
    "empty",
    "unattended",
    "busy",
    "number is busy",
    "number is busy x2",
    "currently busy",
    "na - number currently busy",
    "busy/rejection/na",
    "busy /rejection /na",
    "busy // rejection // na",
    "rejected",
    "rejected,busy",
    "rejected - the number you've dialed is currently busy",
    "rejected-the number you've dialed is currently busy",
    "na rej",
    "na / ringing",
    "db",
    "na db",
    "nadb",
    "dbusy",
    "db x2",
    "cnbr",
    "na cnbr",
    "dvm cnbr",
    "ndt",
    "dnd",
    "switched off",
    "sw off",
    "dvm busy",
    "currently busy",
    "dvm busy",
    "not available",
    "not reachable",
    "on call",
    "not ringing",
    "no answer 5 up",
    "na / busy",
    "na vm*2",
    "na-vm*2",
    "na-vm",
    "nvm",
    "na x2",
    "v2",
    "no ans",
    "no answer beep",
    "na - the number you've dialed is not answering",
    "na - number you've dialed is not answering",
    "ringing no answer",
    "ringing/no answer",
    "ringing / rejected",
    "ringing/rejected",
    "ringing/no answer",
    "ringing/no answer;",
    "ringing no answer;",
    "pu intro hu",
    "na fvm",
    "busy line",
    "direct busy",
    "call failed",
    "network error",
    "unable to accept calls",
    "unable to receive calls",
    "fw to vm",
    "silent beeping",
    "email riverquode - missed call was sent by hamza has",
    "email riverquode - missed call was sent by khaled qa",
    "email riverquode - missed call was sent by taylan bo",
    "email besoin d'aide avec votre compte de trading riverquode? was sent by kossi kp",
    "email fintana - let's activate your ai investment account was sent by tania wi",
    "email how to make a deposit with fintana was sent by alda ga",
    "ai screener na",
    "In Progress",
}
_DATABASE_CHECK_LOG_SHEET_CACHE: str | None = None
_DATABASE_CHECK_NON_ACTION_COMMENTS_NORMALIZED = {
    re.sub(r"[^a-z0-9]+", " ", item.strip().casefold()).strip()
    for item in DATABASE_CHECK_NON_ACTION_COMMENTS
}
_DATABASE_CHECK_NON_ACTION_PATTERNS = [
    re.compile(r"\b(?:na|navm|nadb|nvm|vm|dvm|rej|db|hu|puhu|pu hu)\b"),
    re.compile(r"\bbusy(?: line)?\b|\bdirect busy\b|\bdbusy\b"),
    re.compile(r"\bringing(?: na| no answer| but not answering| but rejecting)?\b"),
    re.compile(r"\b(?:switched off|sw off|call failed|network error)\b"),
    re.compile(r"\bunable to (?:accept|receive) calls\b"),
    re.compile(r"\b(?:fw to vm|voice ?mail|voicemail)\b"),
    re.compile(r"\bno answer(?: \d+| 5 up| vm| beep)?\b|\bno ans\b"),
    re.compile(r"\b(?:silent beeping|beeping|beeps|beep)\b"),
    re.compile(r"\b(?:on call|not reachable|not available|not ringing)\b"),
    re.compile(r"\bin progress\b"),
]


def _read_static_file(filename: str) -> bytes:
    return (PROJECT_ROOT / filename).read_bytes()


def _json_bytes(payload: dict[str, Any]) -> bytes:
    return json.dumps(payload).encode("utf-8")


def _field(form: cgi.FieldStorage, name: str) -> cgi.FieldStorage | None:
    if name not in form:
        return None
    value = form[name]
    if isinstance(value, list):
        return value[0] if value else None
    return value


def _field_text(form: cgi.FieldStorage, name: str) -> str:
    raw = form.getfirst(name, default=None)
    if raw is None:
        value = _field(form, name)
        if value is None:
            return ""
        raw = value.value
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8", errors="replace")
    return str(raw or "").strip()


def _query_text(handler: BaseHTTPRequestHandler, name: str) -> str:
    parsed = urllib.parse.urlparse(handler.path)
    values = urllib.parse.parse_qs(parsed.query, keep_blank_values=True).get(name)
    if not values:
        return ""
    return str(values[0]).strip()


def _header_text(handler: BaseHTTPRequestHandler, name: str) -> str:
    headers = getattr(handler, "headers", None)
    if headers is None:
        return ""

    candidates = {
        name,
        name.title(),
        name.upper(),
        name.lower(),
        name.replace("_", "-"),
        name.replace("_", "-").title(),
    }
    if hasattr(headers, "get"):
        for candidate in candidates:
            value = headers.get(candidate)
            if value:
                return urllib.parse.unquote(str(value).strip())

    if hasattr(headers, "items"):
        target = name.casefold()
        for key, value in headers.items():
            if key.casefold() == target and value:
                return urllib.parse.unquote(str(value).strip())
    return ""


def _field_text_with_query_fallback(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    name: str,
) -> str:
    value = _field_text(form, name)
    if value:
        return value
    header_name = f"x-report-{name.replace('_', '-')}"
    header_value = _header_text(handler, header_name)
    if header_value:
        return header_value
    return _query_text(handler, name)


def _resolve_report_program(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    monthly_comments_upload_id: str,
    upload_meta: dict[str, str] | None = None,
) -> str:
    if monthly_comments_upload_id:
        return PROGRAM_C
    if upload_meta and upload_meta.get("program") in {PROGRAM_A, PROGRAM_B, PROGRAM_C}:
        return upload_meta["program"]
    program = _field_text_with_query_fallback(handler, form, "program") or PROGRAM_A
    if program not in {PROGRAM_A, PROGRAM_B, PROGRAM_C}:
        raise ValueError("Please select a report program.")
    return program


def _resolve_pivot_name(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    upload_meta: dict[str, str] | None = None,
) -> str:
    pivot_name = _field_text_with_query_fallback(handler, form, "pivot_name")
    if pivot_name:
        return pivot_name
    if upload_meta:
        return str(upload_meta.get("pivot_name", "")).strip()
    return ""


def _discover_crm_indices(form: cgi.FieldStorage) -> list[int]:
    indices: set[int] = set()
    for key in form.keys():
        match = re.fullmatch(r"crm_file_(\d+)", str(key))
        if match:
            indices.add(int(match.group(1)))
    return sorted(indices)


def _collect_crm_uploads(
    form: cgi.FieldStorage,
    directory: Path,
) -> tuple[list[Path], list[str]]:
    indices = _discover_crm_indices(form)
    if not indices:
        raise ValueError("At least one CRM file is required.")

    crm_files: list[Path] = []
    platforms: list[str] = []
    for index in indices:
        crm_field = _field(form, f"crm_file_{index}")
        platform = _field_text(form, f"platform_{index}")
        has_upload = _has_upload(crm_field)
        if not has_upload and not platform:
            continue
        if not has_upload:
            raise ValueError(f"Please upload CRM file #{index + 1}.")
        if not platform:
            raise ValueError(
                f"Platform name for CRM file #{index + 1} is required."
            )
        crm_path = _save_upload(
            crm_field,
            directory,
            f"CRM file #{index + 1}",
        )
        crm_files.append(crm_path)
        platforms.append(platform)

    if not crm_files:
        raise ValueError("At least one CRM file is required.")
    return crm_files, platforms


def _resolve_monthly_comments_upload_id(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
) -> str:
    return _field_text_with_query_fallback(
        handler, form, "monthly_comments_upload_id"
    )


def _resolve_monthly_comments_path(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    directory: Path,
    upload_id: str,
) -> Path:
    if upload_id:
        try:
            return resolve_uploaded_file(upload_id)
        except ValueError as exc:
            raise ValueError(
                "Monthly comments upload was not found on the server. "
                "Please refresh the page, re-select the monthly comments file, "
                "and try again."
            ) from exc

    monthly_field = _field(form, "monthly_comments_report")
    if _has_upload(monthly_field):
        return _save_upload(
            monthly_field,
            directory,
            "Monthly comments report",
        )

    raise ValueError("Please upload the monthly comments .xlsx file.")


def _resolve_toggle(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    name: str,
    *,
    default: bool = False,
) -> bool:
    value = _field_text_with_query_fallback(handler, form, name)
    if not value:
        return default
    return _is_truthy(value)


def _field_int_with_fallback(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
    name: str,
) -> int:
    value = _field_text_with_query_fallback(handler, form, name)
    if not value:
        raise ValueError(f"{name} is required.")
    return int(value)


def _handle_monthly_chunk_upload(
    handler: BaseHTTPRequestHandler,
    form: cgi.FieldStorage,
) -> bytes:
    upload_id = _field_text_with_query_fallback(handler, form, "upload_id")
    chunk_index = _field_int_with_fallback(handler, form, "chunk_index")
    total_chunks = _field_int_with_fallback(handler, form, "total_chunks")
    filename = (
        _field_text_with_query_fallback(handler, form, "filename")
        or "monthly_comments.xlsx"
    )
    chunk_field = _field(form, "chunk")
    if chunk_field is None or not getattr(chunk_field, "file", None):
        raise ValueError("Upload chunk file is required.")

    chunk_bytes = chunk_field.file.read()
    if not chunk_bytes:
        raise ValueError("Upload chunk file is empty.")
    if len(chunk_bytes) > CHUNK_UPLOAD_MAX_BYTES:
        raise ValueError(
            f"Each upload chunk must be {CHUNK_UPLOAD_MAX_BYTES // (1024 * 1024)} MB or smaller."
        )

    output_path = save_chunk(
        upload_id,
        chunk_index,
        total_chunks,
        filename,
        chunk_bytes,
        pivot_name=_field_text_with_query_fallback(handler, form, "pivot_name") or None,
        program=_field_text_with_query_fallback(handler, form, "program") or None,
        crm_count=_field_text_with_query_fallback(handler, form, "crm_count") or None,
    )
    return _json_bytes(
        {
            "ok": True,
            "upload_id": upload_id,
            "chunk_index": chunk_index,
            "total_chunks": total_chunks,
            "complete": output_path is not None,
        }
    )


def _optional_text(form: cgi.FieldStorage, name: str) -> str | None:
    value = _field_text(form, name)
    return value or None


def _is_truthy(value: str) -> bool:
    return value.strip().casefold() in {"1", "true", "yes", "on"}


def _safe_filename(raw_name: str | None, fallback: str) -> str:
    name = Path(raw_name or fallback).name
    name = re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("._")
    return name or fallback


def _output_filename(raw_name: str, fallback: str) -> str:
    name = _safe_filename(raw_name, fallback)
    if not name.lower().endswith(".xlsx"):
        name = f"{name}.xlsx"
    return name


def _app_from_form(form: cgi.FieldStorage) -> str:
    app = _field_text(form, "app") or APP_REPORT
    if app not in {APP_REPORT, APP_LEAD_SPLITTER, APP_CR, APP_DATABASE_CHECK}:
        raise ValueError("Please select Report, Lead Splitter, CR, or Database check.")
    return app


def _program_from_form(form: cgi.FieldStorage) -> str:
    program = _field_text(form, "program") or PROGRAM_A
    if program not in {PROGRAM_A, PROGRAM_B, PROGRAM_C}:
        raise ValueError("Please select a report program.")
    return program


def _save_upload(
    field: cgi.FieldStorage | None,
    directory: Path,
    fallback_name: str,
) -> Path:
    if field is None or not field.filename:
        raise ValueError(f"Please upload {fallback_name}.")

    filename = _safe_filename(field.filename, fallback_name)
    path = directory / filename
    data = field.file.read()
    if not data:
        raise ValueError(f"{filename} is empty.")
    path.write_bytes(data)
    return path


def _has_upload(field: cgi.FieldStorage | None) -> bool:
    return bool(field is not None and field.filename)


def _zip_files(file_paths: list[Path]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file_path in file_paths:
            archive.write(file_path, arcname=file_path.name)
    return buffer.getvalue()


def _get_google_private_key() -> str:
    for key in ("gmail", "GMAIL", "GOOGLE_PRIVATE_KEY"):
        raw = os.environ.get(key)
        if raw:
            return raw.strip().replace("\\n", "\n")
    raise ValueError(
        "Missing Google private key. Set environment variable 'gmail' with the service account private key."
    )


def _build_sheets_service(scopes: list[str]) -> Any:
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build as build_google_service
    except ImportError as exc:
        raise ValueError(
            "Google Sheets dependencies are missing. Install google-api-python-client and google-auth."
        ) from exc

    service_account_info = {
        "type": "service_account",
        "client_email": DATABASE_CHECK_LOG_SERVICE_ACCOUNT_EMAIL,
        "private_key": _get_google_private_key(),
        "token_uri": "https://oauth2.googleapis.com/token",
    }
    credentials = service_account.Credentials.from_service_account_info(
        service_account_info,
        scopes=scopes,
    )
    return build_google_service("sheets", "v4", credentials=credentials, cache_discovery=False)


def _quote_sheet_name(sheet_name: str) -> str:
    return "'" + sheet_name.replace("'", "''") + "'"


def _parse_updated_row_number(updated_range: str) -> int | None:
    if not updated_range:
        return None
    match = re.search(r"![A-Z]+(\d+)(?::[A-Z]+\d+)?$", updated_range)
    if not match:
        return None
    return int(match.group(1))


def _resolve_database_check_log_sheet_name(service: Any) -> str:
    global _DATABASE_CHECK_LOG_SHEET_CACHE
    if _DATABASE_CHECK_LOG_SHEET_CACHE:
        return _DATABASE_CHECK_LOG_SHEET_CACHE

    metadata = service.spreadsheets().get(
        spreadsheetId=DATABASE_CHECK_LOG_SPREADSHEET_ID,
        fields="sheets(properties(title))",
    ).execute()
    titles = [
        sheet.get("properties", {}).get("title", "")
        for sheet in metadata.get("sheets", [])
        if sheet.get("properties", {}).get("title")
    ]
    if not titles:
        raise ValueError("Database-check log spreadsheet has no sheets.")

    if DATABASE_CHECK_LOG_SHEET_NAME and DATABASE_CHECK_LOG_SHEET_NAME in titles:
        _DATABASE_CHECK_LOG_SHEET_CACHE = DATABASE_CHECK_LOG_SHEET_NAME
        return _DATABASE_CHECK_LOG_SHEET_CACHE

    required_headers = {"username", "dateandtime", "outputs"}
    for title in titles:
        header_response = service.spreadsheets().values().get(
            spreadsheetId=DATABASE_CHECK_LOG_SPREADSHEET_ID,
            range=f"{_quote_sheet_name(title)}!1:1",
        ).execute()
        header_values = header_response.get("values") or [[]]
        header_row = header_values[0] if header_values else []
        normalized_headers = {_normalize_header_key(value) for value in header_row}
        if required_headers.issubset(normalized_headers):
            _DATABASE_CHECK_LOG_SHEET_CACHE = title
            return _DATABASE_CHECK_LOG_SHEET_CACHE

    raise ValueError(
        "Could not find a log sheet with headers: username, Date and time, outputs."
    )


def _append_database_check_login(username: str) -> dict[str, Any]:
    if not DATABASE_CHECK_LOG_SPREADSHEET_ID:
        raise ValueError("Database-check log spreadsheet ID is not configured.")

    service = _build_sheets_service(["https://www.googleapis.com/auth/spreadsheets"])
    sheet_name = _resolve_database_check_log_sheet_name(service)
    login_time = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    append_response = service.spreadsheets().values().append(
        spreadsheetId=DATABASE_CHECK_LOG_SPREADSHEET_ID,
        range=f"{_quote_sheet_name(sheet_name)}!A:C",
        valueInputOption="USER_ENTERED",
        insertDataOption="INSERT_ROWS",
        body={"values": [[username, login_time, 0]]},
    ).execute()

    updated_range = append_response.get("updates", {}).get("updatedRange", "")
    row_number = _parse_updated_row_number(updated_range)
    if row_number is None:
        raise ValueError("Could not determine login row in the Database-check log sheet.")
    return {"username": username, "login_time": login_time, "log_row": row_number}


def _increment_database_check_output(username: str, log_row: int) -> int:
    service = _build_sheets_service(["https://www.googleapis.com/auth/spreadsheets"])
    sheet_name = _resolve_database_check_log_sheet_name(service)
    row_range = f"{_quote_sheet_name(sheet_name)}!A{log_row}:C{log_row}"
    row_response = service.spreadsheets().values().get(
        spreadsheetId=DATABASE_CHECK_LOG_SPREADSHEET_ID,
        range=row_range,
    ).execute()
    row_values = row_response.get("values", [])
    if not row_values:
        raise ValueError("Database check session was not found. Please log in again.")

    row = row_values[0]
    logged_username = str(row[0]).strip() if row else ""
    if not logged_username or logged_username.casefold() != username.casefold():
        raise ValueError("Database check session is invalid. Please log in again.")

    current_outputs_raw = row[2] if len(row) >= 3 else 0
    try:
        current_outputs = int(float(str(current_outputs_raw).strip() or "0"))
    except ValueError:
        current_outputs = 0
    next_outputs = current_outputs + 1

    service.spreadsheets().values().update(
        spreadsheetId=DATABASE_CHECK_LOG_SPREADSHEET_ID,
        range=f"{_quote_sheet_name(sheet_name)}!C{log_row}",
        valueInputOption="USER_ENTERED",
        body={"values": [[next_outputs]]},
    ).execute()
    return next_outputs


def _database_check_login(form: cgi.FieldStorage) -> dict[str, Any]:
    username = _field_text(form, "database_username")
    password = _field_text(form, "database_password")

    if not username:
        raise ValueError("Username is required.")
    if password != DATABASE_CHECK_PASSWORD:
        raise ValueError("Incorrect password.")

    login_row = _append_database_check_login(username)
    return {"ok": True, **login_row}


def _database_check_validate_session_and_count_output(form: cgi.FieldStorage) -> None:
    username = _field_text(form, "database_username")
    row_text = _field_text(form, "database_log_row")
    if not username or not row_text:
        raise ValueError("Please log in to Database check first.")
    try:
        log_row = int(row_text)
    except ValueError as exc:
        raise ValueError("Database check session is invalid. Please log in again.") from exc
    if log_row < 1:
        raise ValueError("Database check session is invalid. Please log in again.")

    _increment_database_check_output(username, log_row)


def _encode_multipart_bytes(
    *,
    filename: str,
    file_bytes: bytes,
    mime_type: str,
    field_name: str = "file",
    text_fields: dict[str, str] | None = None,
) -> tuple[bytes, str]:
    boundary = f"----CursorBoundary{uuid.uuid4().hex}"
    safe_name = _safe_filename(filename, "upload.bin")

    body_chunks: list[bytes] = []
    for key, value in (text_fields or {}).items():
        field_name_safe = str(key).strip()
        if not field_name_safe:
            continue
        body_chunks.append(
            (
                f"--{boundary}\r\n"
                f'Content-Disposition: form-data; name="{field_name_safe}"\r\n\r\n'
                f"{value}\r\n"
            ).encode("utf-8")
        )

    body_chunks.append(
        (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{field_name}"; filename="{safe_name}"\r\n'
            f"Content-Type: {mime_type}\r\n\r\n"
        ).encode("utf-8")
        + file_bytes
        + b"\r\n"
    )
    body_chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
    body = b"".join(body_chunks)

    return body, f"multipart/form-data; boundary={boundary}"


def _encode_multipart_upload(file_path: Path, field_name: str = "file") -> tuple[bytes, str]:
    filename = _safe_filename(file_path.name, "upload.xlsx")
    mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    file_bytes = file_path.read_bytes()
    return _encode_multipart_bytes(
        filename=filename,
        file_bytes=file_bytes,
        mime_type=mime_type,
        field_name=field_name,
    )


def _normalize_header_key(value: Any) -> str:
    text = str(value or "").strip().casefold()
    return re.sub(r"[^a-z0-9]+", "", text)


def _normalize_match_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).casefold()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip().casefold()

    text = str(value).strip()
    if not text:
        return ""
    numeric_match = re.fullmatch(r"-?\d+(?:\.0+)?", text)
    if numeric_match:
        return str(int(float(text)))
    return text.casefold()


def _normalize_comment_for_matching(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().casefold()).strip()


def _split_comment_entries(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []

    # Some exports store escaped newlines as literal "\n" sequences.
    text = (
        text.replace("\\r\\n", "\n")
        .replace("\\n", "\n")
        .replace("\\r", "\n")
    )
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    entries: list[str] = []
    for line in lines:
        parts = [part.strip() for part in line.split(";") if part.strip()]
        if len(parts) <= 1:
            entries.append(line)
            continue
        for part in parts:
            entries.append(f"{part};")
    return entries


def _format_database_comment_entry(entry: str, row_agent: Any) -> str:
    text = str(entry or "").strip()
    if not text:
        return ""

    agent_fallback = str(row_agent or "").strip()
    pipe_match = re.match(
        r"^(?P<ts>\d{4}-\d{2}-\d{2}\s+\d{1,2}:\d{2})\s*\|\s*(?P<agent>[^|]+?)\s*\|\s*(?P<comment>.*)$",
        text,
    )
    if pipe_match:
        ts = pipe_match.group("ts").strip()
        agent_name = pipe_match.group("agent").strip() or agent_fallback
        comment = pipe_match.group("comment").strip()
        if agent_name:
            return f"|| {ts} | {agent_name} | {comment}"
        return f"|| {ts} | {comment}"

    dash_match = re.match(
        r"^(?P<ts>\d{4}-\d{2}-\d{2}\s+\d{1,2}:\d{2})\s*-\s*(?P<comment>.*)$",
        text,
    )
    if dash_match:
        ts = dash_match.group("ts").strip()
        comment = dash_match.group("comment").strip()
        if agent_fallback:
            return f"|| {ts} | {agent_fallback} | {comment}"
        return f"|| {ts} | {comment}"

    return text


def _comment_payload_text(entry: str) -> str:
    text = str(entry or "").strip()
    if not text:
        return ""
    if " - " in text:
        text = text.split(" - ", 1)[1]
    return text.strip().strip(";").strip()


def _is_non_action_comment(entry: str) -> bool:
    payload = _comment_payload_text(entry)
    normalized = _normalize_comment_for_matching(payload)
    if not normalized:
        return True
    if normalized in _DATABASE_CHECK_NON_ACTION_COMMENTS_NORMALIZED:
        return True
    if any(pattern.search(normalized) for pattern in _DATABASE_CHECK_NON_ACTION_PATTERNS):
        return True
    if re.fullmatch(r"no answer(?: \d+)?", normalized):
        return True
    return False


def _format_database_comments(value: Any, row_agent: Any) -> str:
    entries = _split_comment_entries(value)
    if not entries:
        return ""

    formatted_entries = [_format_database_comment_entry(entry, row_agent) for entry in entries]
    formatted_entries = [entry for entry in formatted_entries if entry]
    return "\n".join(formatted_entries)


def _item_value(item: dict[str, Any], candidate_keys: list[str]) -> Any:
    normalized_map = {_normalize_header_key(key): value for key, value in item.items()}
    for candidate in candidate_keys:
        key = _normalize_header_key(candidate)
        if key in normalized_map:
            return normalized_map[key]
    return ""


def _extract_identity_values(item: dict[str, Any]) -> tuple[Any, Any]:
    cid = _item_value(
        item,
        [
            "CID",
            "cid",
            "Account No",
            "AccountNo",
            "account no",
            "account_no",
            "Account Number",
            "accountnumber",
            "Customer ID",
            "CustomerId",
            "customer_id",
        ],
    )
    brand = _item_value(item, ["Brand", "brand", "Brand Name", "BrandName", "brand_name"])

    if cid and brand:
        return cid, brand

    for key, value in item.items():
        normalized_key = _normalize_header_key(key)
        if not cid and (
            "cid" in normalized_key
            or normalized_key in {"accountno", "accountnumber", "customerid"}
        ):
            cid = value
        if not brand and "brand" in normalized_key:
            brand = value
        if cid and brand:
            break
    return cid, brand


def _is_identity_scalar(value: Any) -> bool:
    return not isinstance(value, (list, tuple, dict, set))


def _looks_like_identity_header(value: Any) -> bool:
    normalized = _normalize_header_key(value)
    return bool(
        normalized
        and (
            "cid" in normalized
            or normalized in {"accountno", "accountnumber", "customerid"}
            or "brand" in normalized
        )
    )


def _records_from_table_rows(rows: list[Any]) -> list[dict[str, Any]]:
    if len(rows) < 2:
        return []
    if not all(isinstance(row, (list, tuple)) for row in rows):
        return []

    header_row = list(rows[0])
    if not any(_looks_like_identity_header(cell) for cell in header_row):
        return []

    headers = [str(cell or "").strip() for cell in header_row]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        row_values = list(row)
        if not any(value not in (None, "") for value in row_values):
            continue
        record = {
            headers[index]: row_values[index] if index < len(row_values) else ""
            for index in range(len(headers))
            if headers[index]
        }
        if record:
            records.append(record)
    return records


def _records_from_columnar_dict(node: dict[str, Any]) -> list[dict[str, Any]]:
    cid_key = None
    brand_key = None
    for key in node.keys():
        normalized = _normalize_header_key(key)
        if cid_key is None and (
            "cid" in normalized
            or normalized in {"accountno", "accountnumber", "customerid"}
        ):
            cid_key = key
        if brand_key is None and "brand" in normalized:
            brand_key = key
    if cid_key is None or brand_key is None:
        return []

    cid_values = node.get(cid_key)
    brand_values = node.get(brand_key)
    if not isinstance(cid_values, list) or not isinstance(brand_values, list):
        return []

    row_count = min(len(cid_values), len(brand_values))
    if row_count == 0:
        return []

    records: list[dict[str, Any]] = []
    for row_index in range(row_count):
        record: dict[str, Any] = {}
        for key, value in node.items():
            if isinstance(value, list):
                record[key] = value[row_index] if row_index < len(value) else ""
            else:
                record[key] = value
        records.append(record)
    return records


def _record_from_text(text: str) -> dict[str, Any] | None:
    if not isinstance(text, str):
        return None
    normalized = text.strip()
    if not normalized:
        return None

    cid_match = re.search(
        r"\b(?:cid|account\s*no|accountnumber|customer\s*id)\b\s*[:=]\s*['\"]?([A-Za-z0-9._-]+)",
        normalized,
        re.IGNORECASE,
    )
    brand_match = re.search(
        r"\b(?:brand|brand\s*name)\b\s*[:=]\s*['\"]?([A-Za-z0-9 _.-]+)",
        normalized,
        re.IGNORECASE,
    )
    if not cid_match or not brand_match:
        return None

    suggested_match = re.search(
        r"\b(?:suggested\s*status|recommended\s*status|recommendationstatus)\b\s*[:=]\s*['\"]?([^,\n\r;]+)",
        normalized,
        re.IGNORECASE,
    )
    reason_match = re.search(
        r"\b(?:reason|explanation|rationale)\b\s*[:=]\s*['\"]?([^\n\r]+)",
        normalized,
        re.IGNORECASE,
    )
    return {
        "CID": cid_match.group(1).strip(),
        "Brand": brand_match.group(1).strip(),
        "Suggested Status": suggested_match.group(1).strip() if suggested_match else "",
        "Reason": reason_match.group(1).strip() if reason_match else "",
    }


def _extract_webhook_data_items(payload: Any) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []

    def _try_parse_json_text(text: str) -> Any:
        stripped = text.strip()
        if not stripped:
            return text

        try:
            return json.loads(stripped)
        except json.JSONDecodeError:
            pass

        try:
            return ast.literal_eval(stripped)
        except (ValueError, SyntaxError):
            pass

        return text

    def parse_json_like(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        text = value.strip()
        if not text:
            return value

        # n8n or LLM responses may wrap JSON in markdown code fences.
        fenced_match = re.search(r"```(?:json)?\s*([\s\S]*?)```", text, re.IGNORECASE)
        if fenced_match:
            text = fenced_match.group(1).strip()

        parsed: Any = text
        for _ in range(3):
            next_value = _try_parse_json_text(parsed) if isinstance(parsed, str) else parsed
            if next_value is parsed:
                break
            parsed = next_value

        if not isinstance(parsed, str):
            return parsed

        # Last attempt: extract a JSON-looking object/array embedded in plain text.
        embedded_match = re.search(r"(\{[\s\S]*\}|\[[\s\S]*\])", parsed)
        if embedded_match:
            embedded = embedded_match.group(1).strip()
            embedded_parsed = _try_parse_json_text(embedded)
            if not isinstance(embedded_parsed, str):
                return embedded_parsed

        # Fallback for CSV-like blocks where first row contains headers.
        if "\n" in text and "," in text:
            try:
                csv_rows = list(csv.reader(text.splitlines()))
            except Exception:
                csv_rows = []
            table_records = _records_from_table_rows(csv_rows)
            if table_records:
                return table_records

        return value

    def walk(node: Any) -> None:
        node = parse_json_like(node)

        if isinstance(node, list):
            table_records = _records_from_table_rows(node)
            if table_records:
                for record in table_records:
                    walk(record)
            for entry in node:
                walk(entry)
            return

        if not isinstance(node, dict):
            if isinstance(node, str):
                record = _record_from_text(node)
                if record:
                    items.append(record)
            return

        cid, brand = _extract_identity_values(node)
        if (
            _is_identity_scalar(cid)
            and _is_identity_scalar(brand)
            and _normalize_match_value(cid)
            and _normalize_match_value(brand)
        ):
            items.append(node)

        for record in _records_from_columnar_dict(node):
            walk(record)

        data_payload = None
        for key, value in node.items():
            if _normalize_header_key(key) == "data":
                data_payload = value
                break

        if isinstance(data_payload, (list, dict, str)):
            walk(data_payload)

        for value in node.values():
            if value is data_payload:
                continue
            if isinstance(value, (list, dict, str)):
                walk(value)

    walk(payload)
    return items


def _candidate_payload_nodes(payload: Any) -> list[Any]:
    candidates: list[Any] = []
    seen_ids: set[int] = set()

    def add(node: Any) -> None:
        if node is None:
            return
        node_id = id(node)
        if node_id in seen_ids:
            return
        seen_ids.add(node_id)
        candidates.append(node)

    add(payload)

    queue: list[Any] = [payload]
    while queue:
        node = queue.pop(0)

        if isinstance(node, list):
            if node:
                add(node[0])  # Explicit File[0] / payload[0] handling.
                queue.append(node[0])
            continue

        if not isinstance(node, dict):
            continue

        for key, value in node.items():
            normalized_key = _normalize_header_key(key)
            if normalized_key in {
                "file",
                "data",
                "items",
                "records",
                "result",
                "results",
                "output",
                "response",
            }:
                add(value)
                queue.append(value)
                if isinstance(value, list) and value:
                    add(value[0])  # Explicitly support key-based list first item lookup.
                    queue.append(value[0])

    return candidates


def _build_database_suggestions(payload: Any) -> dict[tuple[str, str], tuple[Any, Any]]:
    items: list[dict[str, Any]] = []
    for candidate in _candidate_payload_nodes(payload):
        items.extend(_extract_webhook_data_items(candidate))
    suggestions: dict[tuple[str, str], tuple[Any, Any]] = {}

    for item in items:
        cid, brand = _extract_identity_values(item)
        if not (_is_identity_scalar(cid) and _is_identity_scalar(brand)):
            continue
        cid_key = _normalize_match_value(cid)
        brand_key = _normalize_match_value(brand)
        if not cid_key or not brand_key:
            continue

        suggested_status = _item_value(
            item,
            [
                "Suggested status",
                "Suggested Status",
                "suggested_status",
                "SuggestedStatus",
                "RecommendationStatus",
                "RecommendedStatus",
            ],
        )
        reason = _item_value(item, ["Reason", "reason", "Explanation", "Rationale"])
        key = (cid_key, brand_key)

        existing = suggestions.get(key)
        if existing is None:
            suggestions[key] = (suggested_status, reason)
        else:
            current_status, current_reason = existing
            suggestions[key] = (
                current_status if current_status else suggested_status,
                current_reason if current_reason else reason,
            )

    if not suggestions:
        payload_error = _extract_payload_error_message(payload)
        if payload_error:
            raise ValueError(payload_error)
        raise ValueError(
            "Webhook JSON did not include usable Data items with CID and brand fields."
        )
    return suggestions


def _extract_payload_error_message(payload: Any) -> str | None:
    candidates: list[str] = []

    def walk(node: Any) -> None:
        if isinstance(node, list):
            for entry in node:
                walk(entry)
            return
        if not isinstance(node, dict):
            return

        cid, brand = _extract_identity_values(node)
        if _normalize_match_value(cid) and _normalize_match_value(brand):
            return

        message = node.get("message")
        error = node.get("error")
        hint = node.get("hint")
        code = node.get("code")
        status = node.get("status")

        if isinstance(error, str) and error.strip():
            candidates.append(error.strip())
        if isinstance(message, str) and message.strip():
            text = message.strip()
            prefix = f"{code}" if code not in (None, "") else f"{status}" if status not in (None, "") else ""
            candidates.append(f"{prefix}: {text}" if prefix else text)
        if isinstance(hint, str) and hint.strip():
            candidates.append(f"Hint: {hint.strip()}")

        for value in node.values():
            if isinstance(value, (list, dict)):
                walk(value)

    for candidate in _candidate_payload_nodes(payload):
        walk(candidate)
    if not candidates:
        return None
    return "Webhook response error: " + " | ".join(dict.fromkeys(candidates))


def _resolve_input_header_indexes(worksheet) -> tuple[int, dict[str, int]]:
    required_keys = {column: _normalize_header_key(column) for column in DATABASE_CHECK_INPUT_COLUMNS}
    best_row_index = 0
    best_count = -1
    best_index_map: dict[str, int] = {}

    for row_index in (1, 2, 3):
        row_values = next(
            worksheet.iter_rows(min_row=row_index, max_row=row_index, values_only=True),
            (),
        )
        normalized_to_index: dict[str, int] = {}
        for col_index, value in enumerate(row_values):
            normalized = _normalize_header_key(value)
            if normalized and normalized not in normalized_to_index:
                normalized_to_index[normalized] = col_index

        found_labels = {
            label: normalized_to_index[key]
            for label, key in required_keys.items()
            if key in normalized_to_index
        }
        if len(found_labels) > best_count:
            best_count = len(found_labels)
            best_row_index = row_index
            best_index_map = found_labels
        if len(found_labels) == len(DATABASE_CHECK_INPUT_COLUMNS):
            return row_index, found_labels

    missing = [label for label in DATABASE_CHECK_INPUT_COLUMNS if label not in best_index_map]
    raise ValueError(
        "Could not find required headers in rows 1 to 3. Missing columns: "
        + ", ".join(missing)
    )


def _build_database_check_output(input_path: Path, payload: Any, output_path: Path) -> None:
    suggestions = _build_database_suggestions(payload)

    workbook = load_workbook(input_path, data_only=True)
    worksheet = workbook.active
    header_row_index, header_indexes = _resolve_input_header_indexes(worksheet)

    output_workbook = Workbook()
    output_sheet = output_workbook.active
    output_sheet.title = "Database check"
    output_sheet.append(DATABASE_CHECK_OUTPUT_COLUMNS)
    cleaned_input_sheet = output_workbook.create_sheet("Input cleanup")
    cleaned_input_sheet.append(DATABASE_CHECK_INPUT_COLUMNS)

    matched_count = 0
    for row in worksheet.iter_rows(min_row=header_row_index + 1, values_only=True):
        if not row or all(value in (None, "") for value in row):
            continue
        cleaned_row = [row[header_indexes[column]] for column in DATABASE_CHECK_INPUT_COLUMNS]
        cleaned_input_sheet.append(cleaned_row)

        account_no = row[header_indexes["Account No"]]
        brand = row[header_indexes["Brand"]]
        match_key = (_normalize_match_value(account_no), _normalize_match_value(brand))
        if not all(match_key) or match_key not in suggestions:
            continue

        suggested_status, reason = suggestions[match_key]
        output_row = list(cleaned_row)
        output_row.extend([suggested_status, reason])
        output_sheet.append(output_row)
        matched_count += 1

    if matched_count == 0:
        raise ValueError("No rows matched CID + brand between webhook Data and input file.")

    for col_idx, header in enumerate(DATABASE_CHECK_OUTPUT_COLUMNS, start=1):
        column_values = [
            str(output_sheet.cell(r, col_idx).value or "")
            for r in range(1, min(output_sheet.max_row, 300) + 1)
        ]
        max_len = max([len(str(header)), *(len(value) for value in column_values)])
        output_sheet.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 48)

    for col_idx, header in enumerate(DATABASE_CHECK_INPUT_COLUMNS, start=1):
        column_values = [
            str(cleaned_input_sheet.cell(r, col_idx).value or "")
            for r in range(1, min(cleaned_input_sheet.max_row, 300) + 1)
        ]
        max_len = max([len(str(header)), *(len(value) for value in column_values)])
        cleaned_input_sheet.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 48)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_workbook.save(output_path)


def _build_database_check_webhook_records(input_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, data_only=True)
    worksheet = workbook.active
    header_row_index, header_indexes = _resolve_input_header_indexes(worksheet)

    records: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=header_row_index + 1, values_only=True):
        if not row or all(value in (None, "") for value in row):
            continue

        raw_comments = row[header_indexes["Last 10 Comments"]]
        row_agent = row[header_indexes["Current Assigned Agent"]]
        formatted_comments = _format_database_comments(raw_comments, row_agent)

        record: dict[str, Any] = {}
        for column in DATABASE_CHECK_WEBHOOK_COLUMNS:
            payload_key = DATABASE_CHECK_WEBHOOK_KEY_BY_COLUMN[column]
            if column == "Last 10 Comments":
                record[payload_key] = formatted_comments
            else:
                record[payload_key] = row[header_indexes[column]]

        if all(value in (None, "") for value in record.values()):
            continue
        records.append(record)

    if not records:
        raise ValueError("No usable rows were found to build the Database-check JSON payload.")
    return records


def _build_database_check_webhook_upload(
    input_path: Path,
    callback_url: str = "",
) -> tuple[bytes, str]:
    records = _build_database_check_webhook_records(input_path)
    payload_bytes = json.dumps({"data": records}, ensure_ascii=False).encode("utf-8")
    text_fields: dict[str, str] = {}
    if callback_url:
        text_fields = {
            "callback_url": callback_url,
            "callbackUrl": callback_url,
        }
    return _encode_multipart_bytes(
        filename=f"{input_path.stem}_database_check_payload.json",
        file_bytes=payload_bytes,
        mime_type="application/json",
        field_name="file",
        text_fields=text_fields,
    )


def _callback_url_for_request(handler: BaseHTTPRequestHandler) -> str:
    if DATABASE_CHECK_CALLBACK_URL:
        return DATABASE_CHECK_CALLBACK_URL

    host = (
        handler.headers.get("x-forwarded-host")
        or handler.headers.get("host")
        or ""
    ).strip()
    if not host:
        return ""

    proto = (handler.headers.get("x-forwarded-proto") or "https").strip() or "https"
    path = DATABASE_CHECK_CALLBACK_PATH
    if not path.startswith("/"):
        path = "/" + path
    return f"{proto}://{host}{path}"


def _database_check_timeout_seconds() -> float | None:
    if not DATABASE_CHECK_TIMEOUT_SECONDS_RAW:
        return None
    try:
        timeout_seconds = float(DATABASE_CHECK_TIMEOUT_SECONDS_RAW)
    except ValueError:
        return None
    if timeout_seconds <= 0:
        return None
    return timeout_seconds


def _request_webhook_json(file_path: Path, callback_url: str = "") -> Any:
    if not DATABASE_CHECK_WEBHOOK_URL:
        raise ValueError("Database-check webhook URL is not configured.")

    body, content_type = _build_database_check_webhook_upload(
        file_path,
        callback_url=callback_url,
    )
    request = urllib.request.Request(
        DATABASE_CHECK_WEBHOOK_URL,
        data=body,
        method="POST",
        headers={"Content-Type": content_type, "Accept": "*/*"},
    )
    timeout_seconds = _database_check_timeout_seconds()

    try:
        if timeout_seconds is None:
            response_context = urllib.request.urlopen(request)
        else:
            response_context = urllib.request.urlopen(request, timeout=timeout_seconds)
        with response_context as response:
            response_bytes = response.read()
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace").strip()
        message = details or str(exc.reason)
        raise ValueError(f"Webhook request failed ({exc.code}): {message}") from exc
    except urllib.error.URLError as exc:
        raise ValueError(f"Webhook request failed: {exc.reason}") from exc

    if not response_bytes:
        raise ValueError("Webhook returned an empty response.")

    decoded_text = response_bytes.decode("utf-8", errors="replace")
    try:
        payload = json.loads(decoded_text)
    except json.JSONDecodeError:
        try:
            payload = ast.literal_eval(decoded_text)
        except (ValueError, SyntaxError):
            payload = decoded_text

    if isinstance(payload, dict):
        error_message = payload.get("error") or payload.get("message")
        if error_message:
            raise ValueError(str(error_message))
    return payload


def _parse_form(handler: BaseHTTPRequestHandler) -> cgi.FieldStorage:
    content_type = handler.headers.get("content-type", "")
    if not content_type.lower().startswith("multipart/form-data"):
        raise ValueError("Upload form must use multipart/form-data.")

    content_length = int(handler.headers.get("content-length", "0") or "0")
    if content_length <= 0:
        raise ValueError("No upload data was received.")
    if content_length > MAX_UPLOAD_BYTES:
        raise ValueError("Upload is too large. Please keep the total upload under 45 MB.")

    body = handler.rfile.read(content_length)
    headers = Message()
    headers["content-type"] = content_type
    headers["content-length"] = str(len(body))
    environ = {
        "REQUEST_METHOD": "POST",
        "CONTENT_TYPE": content_type,
        "CONTENT_LENGTH": str(len(body)),
    }
    return cgi.FieldStorage(
        fp=BytesIO(body),
        headers=headers,
        environ=environ,
        keep_blank_values=True,
    )


class handler(BaseHTTPRequestHandler):
    def _send_static_file(self, filename: str, content_type: str) -> None:
        try:
            content = _read_static_file(filename)
        except FileNotFoundError:
            self.send_error(404, "File not found")
            return

        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_GET(self) -> None:
        if self.path in {"/", "/index.html"}:
            self._send_static_file("index.html", "text/html; charset=utf-8")
            return

        if self.path == "/styles.css":
            self._send_static_file("styles.css", "text/css; charset=utf-8")
            return

        self.send_error(404, "Not found")

    def do_OPTIONS(self) -> None:
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header(
            "Access-Control-Allow-Headers",
            "Content-Type, X-Report-Pivot-Name, X-Report-Program, X-Report-Crm-Count, "
            "X-Report-Monthly-Comments-Upload-Id, X-Report-Separate-M-Inhouse, "
            "X-Report-Separate-Department, X-Report-Separate-By-Days, X-Report-Output-File",
        )
        self.end_headers()

    def do_POST(self) -> None:
        try:
            form = _parse_form(self)
            report_action = _field_text_with_query_fallback(self, form, "report_action")
            if report_action == REPORT_ACTION_UPLOAD_MONTHLY_CHUNK:
                response_bytes = _handle_monthly_chunk_upload(self, form)
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Content-Length", str(len(response_bytes)))
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(response_bytes)
                return

            app = _app_from_form(form)
            if app == APP_DATABASE_CHECK:
                database_action = (_field_text(form, "database_action") or "run").casefold()
                if database_action == "login":
                    payload = _database_check_login(form)
                    response_bytes = _json_bytes(payload)
                    response_content_type = "application/json; charset=utf-8"
                    response_filename = ""
                    self.send_response(200)
                    self.send_header("Content-Type", response_content_type)
                    self.send_header("Content-Length", str(len(response_bytes)))
                    self.send_header("Access-Control-Allow-Origin", "*")
                    self.end_headers()
                    self.wfile.write(response_bytes)
                    return
                if database_action not in {"run", "generate", "output"}:
                    raise ValueError("Invalid database-check action.")

            with tempfile.TemporaryDirectory() as tmp:
                tmp_path = Path(tmp)
                if app == APP_REPORT:
                    monthly_comments_upload_id = _resolve_monthly_comments_upload_id(
                        self, form
                    )
                    upload_meta = (
                        read_upload_metadata(monthly_comments_upload_id)
                        if monthly_comments_upload_id
                        else {}
                    )
                    program = _resolve_report_program(
                        self, form, monthly_comments_upload_id, upload_meta
                    )
                    pivot_name = _resolve_pivot_name(self, form, upload_meta)
                    if program in {PROGRAM_A, PROGRAM_C} and not pivot_name:
                        raise ValueError("Pivot table name is required for this report program.")

                    powerbi_path: Path | None = None
                    monthly_comments_path: Path | None = None
                    if program == PROGRAM_C:
                        monthly_comments_path = _resolve_monthly_comments_path(
                            self,
                            form,
                            tmp_path,
                            monthly_comments_upload_id,
                        )
                    else:
                        powerbi_path = _save_upload(
                            _field(form, "powerbi_report"),
                            tmp_path,
                            "PowerBI report",
                        )

                    crm_files, platforms = _collect_crm_uploads(form, tmp_path)

                    default_output = {
                        PROGRAM_B: PROGRAM_B_OUTPUT_FILENAME,
                        PROGRAM_C: PROGRAM_C_OUTPUT_FILENAME,
                    }.get(program, PROGRAM_A_OUTPUT_FILENAME)
                    response_filename = _output_filename(
                        _field_text_with_query_fallback(self, form, "output_file"),
                        default_output,
                    )
                    output_path = tmp_path / response_filename
                    common_args = {
                        "crm_files": crm_files,
                        "platforms": platforms,
                        "output_file": output_path,
                        "powerbi_sheet": _optional_text(form, "powerbi_sheet"),
                        "crm_sheet": _optional_text(form, "crm_sheet"),
                    }
                    if powerbi_path is not None:
                        common_args["powerbi_report"] = powerbi_path
                    if monthly_comments_path is not None:
                        common_args["monthly_comments_report"] = monthly_comments_path
                    separate_department = _resolve_toggle(
                        self, form, "separate_department"
                    )
                    separate_by_days = _resolve_toggle(
                        self, form, "separate_by_days"
                    )
                    if program == PROGRAM_B:
                        generated_outputs = program_b_country_report.build_output(
                            **common_args,
                            separate_department=separate_department,
                            separate_by_days=separate_by_days,
                        )
                        if not generated_outputs or len(generated_outputs) == 1:
                            only_output = (
                                generated_outputs[0]
                                if generated_outputs
                                else output_path
                            )
                            response_filename = only_output.name
                            response_bytes = only_output.read_bytes()
                            response_content_type = XLSX_CONTENT_TYPE
                        else:
                            response_filename = f"{output_path.stem}_reports.zip"
                            response_bytes = _zip_files(generated_outputs)
                            response_content_type = "application/zip"
                    else:
                        separate_m_inhousemedia = _resolve_toggle(
                            self,
                            form,
                            "separate_m_inhouse",
                            default=True,
                        )
                        generated_outputs = program_a_report.build_output_files(
                            **common_args,
                            pivot_name=pivot_name,
                            separate_m_inhousemedia=separate_m_inhousemedia,
                            separate_department=separate_department,
                            separate_by_days=separate_by_days,
                        )
                        if len(generated_outputs) == 1:
                            only_output = generated_outputs[0]
                            response_filename = only_output.name
                            response_bytes = only_output.read_bytes()
                            response_content_type = XLSX_CONTENT_TYPE
                        else:
                            response_filename = f"{output_path.stem}_reports.zip"
                            response_bytes = _zip_files(generated_outputs)
                            response_content_type = "application/zip"
                    if monthly_comments_upload_id:
                        cleanup_upload(monthly_comments_upload_id)
                elif app == APP_LEAD_SPLITTER:
                    lead_input = _save_upload(
                        _field(form, "lead_input"),
                        tmp_path,
                        "Lead splitter input file",
                    )
                    lead_kind = (_field_text(form, "lead_kind") or "lead").lower()
                    if lead_kind not in {"lead", "aff", "countries"}:
                        raise ValueError("Invalid Lead Splitter output type requested.")

                    generated_outputs = lead_splitter.build_outputs(
                        input_path=lead_input,
                        output_dir=tmp_path,
                        generate_lead=(lead_kind == "lead"),
                        generate_aff=(lead_kind == "aff"),
                        generate_countries=(lead_kind == "countries"),
                    )
                    selected_output = generated_outputs.get(lead_kind)
                    if selected_output is None:
                        if lead_kind == "aff":
                            raise ValueError("AFF by Status output was not generated from this input.")
                        if lead_kind == "countries":
                            raise ValueError(
                                "Lead splitter by countries output was not generated from this input."
                            )
                        raise ValueError("Lead Splitter output was not generated from this input.")
                    response_filename = selected_output.name
                    response_bytes = selected_output.read_bytes()
                    response_content_type = XLSX_CONTENT_TYPE
                elif app == APP_DATABASE_CHECK:
                    _database_check_validate_session_and_count_output(form)
                    database_input = _save_upload(
                        _field(form, "database_input"),
                        tmp_path,
                        "Database check input file",
                    )
                    callback_url = _callback_url_for_request(self)
                    webhook_payload = _request_webhook_json(
                        database_input,
                        callback_url=callback_url,
                    )
                    response_filename = f"database_check_{database_input.stem}.xlsx"
                    database_output = tmp_path / response_filename
                    _build_database_check_output(
                        database_input,
                        webhook_payload,
                        database_output,
                    )
                    response_bytes = database_output.read_bytes()
                    response_content_type = XLSX_CONTENT_TYPE
                else:
                    cr_input = _save_upload(
                        _field(form, "cr_input"),
                        tmp_path,
                        "CR input file",
                    )
                    response_filename = cr_maker.default_output_filename()
                    cr_output_path = tmp_path / response_filename
                    cr_maker.process(cr_input, cr_output_path)
                    response_bytes = cr_output_path.read_bytes()
                    response_content_type = XLSX_CONTENT_TYPE

        except Exception as exc:
            self.send_response(400)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(_json_bytes({"error": str(exc)}))
            return

        self.send_response(200)
        self.send_header("Content-Type", response_content_type)
        self.send_header(
            "Content-Disposition",
            f'attachment; filename="{response_filename}"',
        )
        self.send_header("Content-Length", str(len(response_bytes)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(response_bytes)
