from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import cy_ftd
from report_generator import (
    CRM_COLUMNS,
    POWERBI_COLUMNS,
    PROGRAM_A_OUTPUT_COLUMNS,
    build_output_files,
)


CY_FTD_HEADERS = [
    "Brand",  # A
    "Desk",  # B
    "C",
    "D",
    "Account No",  # E
    "F",
    "Campaign",  # G-ish (located by header)
    "Country",
    "Registration Time(UTC)",
    "J",
    "K",
    "L",
    "M",
    "FTD Transaction Owner",
    "FTD Cnt",  # O
]


def _write_cy_ftd_file(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["meta"] * len(CY_FTD_HEADERS))
    worksheet.append(["meta"] * len(CY_FTD_HEADERS))
    worksheet.append(CY_FTD_HEADERS)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def _cy_ftd_row(brand, desk, account_no, campaign, country, reg_time, owner, ftd_cnt):
    return [
        brand,
        desk,
        "",
        "",
        account_no,
        "",
        campaign,
        country,
        reg_time,
        "",
        "",
        "",
        "",
        owner,
        ftd_cnt,
    ]


def _write_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


class CyFtdParsingTests(unittest.TestCase):
    def test_strip_owner_code(self) -> None:
        self.assertEqual(cy_ftd.strip_owner_code("Marco Mo (PW1020)"), "Marco Mo")
        self.assertEqual(cy_ftd.strip_owner_code("Jane Doe"), "Jane Doe")

    def test_read_entries_filters_by_ftd_cnt_and_cy_desk(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "cy.xlsx"
            _write_cy_ftd_file(
                path,
                [
                    _cy_ftd_row("BrandA", "CY1-EN", 111, "Camp A", "Cyprus", "2026-08-01 10:00", "Marco Mo (PW1020)", 1),
                    # FTD Cnt empty -> skipped
                    _cy_ftd_row("BrandA", "CY1-EN", 222, "Camp A", "Cyprus", "2026-08-02 10:00", "Ann (X1)", ""),
                    # Non-CY desk -> skipped
                    _cy_ftd_row("BrandA", "TR2-EN", 333, "Camp A", "Turkey", "2026-08-03 10:00", "Bob (X2)", 1),
                    # CY with multi-part desk -> kept
                    _cy_ftd_row("BrandB", "CY1-EN-TR", 444, "Camp B", "Malta", "2026-08-04 10:00", "Sue", 1),
                ],
            )
            entries = cy_ftd.read_cy_ftd_entries(path)
            self.assertEqual(len(entries), 2)
            self.assertEqual(entries[0]["account_no"], 111)
            self.assertEqual(entries[0]["owner"], "Marco Mo")
            self.assertEqual(entries[1]["account_no"], 444)
            self.assertEqual(entries[1]["desk"], "CY1-EN-TR")

    def test_apply_skips_existing_telemarketing_and_adds_missing(self) -> None:
        rows = [
            {"ID": 111, "Platform": "BrandA", "Status": "Telemarketing"},
            {"ID": 999, "Platform": "BrandA", "Status": "Potential"},
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "cy.xlsx"
            _write_cy_ftd_file(
                path,
                [
                    # Already telemarketing in CRM -> skipped.
                    _cy_ftd_row("BrandA", "CY1-EN", 111, "Camp A", "Cyprus", "2026-08-01", "Marco Mo (PW1020)", 1),
                    # New -> added.
                    _cy_ftd_row("BrandA", "CY2-EN", 555, "Camp B", "Malta", "2026-08-05 09:30", "Sara Ali (PW9)", 1),
                ],
            )
            cy_ftd.apply_cy_ftd_rows(rows, path)

        added = [r for r in rows if r.get("ID") == 555]
        self.assertEqual(len(added), 1)
        row = added[0]
        self.assertEqual(row["Platform"], "BrandA")
        self.assertEqual(row["Department"], "CY2-EN")
        self.assertEqual(row["Assigned to"], "Sara Ali")
        self.assertEqual(row["Country"], "Malta")
        self.assertEqual(row["Created"], "2026-08-05 09:30")
        self.assertEqual(row["Campaign"], "Camp B")
        self.assertEqual(row["Status"], "Telemarketing")
        self.assertEqual(row["Call Attempts"], 1)
        # 111 was already telemarketing -> not duplicated.
        self.assertEqual(len([r for r in rows if r.get("ID") == 111]), 1)


class CyFtdEndToEndTests(unittest.TestCase):
    def test_build_output_files_adds_cy_ftd_rows_and_kyc(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            powerbi = root / "powerbi.xlsx"
            crm = root / "crm.xlsx"
            cy = root / "cy.xlsx"
            output = root / "out.xlsx"

            _write_workbook(powerbi, POWERBI_COLUMNS, [[111, "BrandA", "| x ;", 1]])
            _write_workbook(
                crm,
                [*CRM_COLUMNS, "Date of Birth"],
                [
                    [
                        "Lead", 111, "2026-08-01", "Jane", "Sales", "Potential",
                        "TR", "Camp A", "Sub", "Place", "Agent 1", "1990-01-01",
                    ],
                ],
            )
            _write_cy_ftd_file(
                cy,
                [
                    _cy_ftd_row("BrandA", "CY1-EN", 777, "Camp X", "Cyprus", "2026-08-06 12:00", "Omar K (PW77)", 1),
                ],
            )

            lookup = {("777", "branda"): "KYC OK"}
            build_output_files(
                powerbi_report=powerbi,
                crm_files=[crm],
                platforms=["BrandA"],
                pivot_name="Status Pivot",
                output_file=output,
                separate_m_inhousemedia=False,
                telemarketing_kyc_lookup=lookup,
                cy_ftd_report=cy,
            )

            workbook = load_workbook(output, data_only=False)
            ws = workbook["CRM Output"]
            id_col = PROGRAM_A_OUTPUT_COLUMNS.index("ID") + 1
            status_col = PROGRAM_A_OUTPUT_COLUMNS.index("Status") + 1
            assigned_col = PROGRAM_A_OUTPUT_COLUMNS.index("Assigned to") + 1
            comments_col = PROGRAM_A_OUTPUT_COLUMNS.index("Comments") + 1
            dept_col = PROGRAM_A_OUTPUT_COLUMNS.index("Department") + 1

            found = None
            for row_idx in range(2, ws.max_row + 1):
                if str(ws.cell(row_idx, id_col).value) == "777":
                    found = row_idx
                    break
            self.assertIsNotNone(found, "CY FTD row (ID 777) should be in the output")
            self.assertEqual(ws.cell(found, status_col).value, "Telemarketing")
            self.assertEqual(ws.cell(found, assigned_col).value, "Omar K")
            self.assertEqual(ws.cell(found, dept_col).value, "CY1-EN")
            self.assertEqual(ws.cell(found, comments_col).value, "KYC OK")


if __name__ == "__main__":
    unittest.main()
