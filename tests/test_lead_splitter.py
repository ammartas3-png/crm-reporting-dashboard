from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import lead_splitter


def _write_input_with_header_row_3(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["meta"] * len(headers))
    worksheet.append(["meta"] * len(headers))
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


class LeadSplitterAffCrTests(unittest.TestCase):
    def test_aff_totals_use_strict_ftd_divided_by_leads(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "lead_input.xlsx"
            output_dir = root / "out"

            headers = [
                "Ignore0",
                "Desk",
                "Agent",
                "Ignore3",
                "CID",
                "Status",
                "Campaign",
                "Ignore7",
                "Country",
                "Ignore9",
                "Ignore10",
                "Ignore11",
                "Ignore12",
                "Assigned",
                "FTD",
            ]
            rows = [
                ["", "TR1-EN", "Agent A", "", "CID-1", "No Answer", "Aff A", "", "Switzerland", "", "", "", "", "0", "1"],
            ]
            _write_input_with_header_row_3(input_path, headers, rows)

            outputs = lead_splitter.build_outputs(
                input_path=input_path,
                output_dir=output_dir,
                generate_lead=False,
                generate_aff=True,
            )
            aff_output = outputs["aff"]
            workbook = load_workbook(aff_output, data_only=False)
            sheet = workbook["AFF by Status"]

            def _cr_for_label(label: str, label_column: int, cr_column: int) -> float:
                for row_idx in range(2, sheet.max_row + 1):
                    if str(sheet.cell(row_idx, label_column).value or "").strip() == label:
                        return float(sheet.cell(row_idx, cr_column).value or 0.0)
                self.fail(f"Could not find label '{label}' in column {label_column}.")

            # CH table is written in columns A:G.
            self.assertEqual(_cr_for_label("Aff A Total", 3, 7), 0.0)
            self.assertEqual(_cr_for_label("TR Total", 2, 7), 0.0)
            self.assertEqual(_cr_for_label("CH Total", 1, 7), 0.0)


class LeadSplitterCountriesTests(unittest.TestCase):
    def test_countries_output_exists_and_uses_aff_cr_rules(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "lead_input.xlsx"
            output_dir = root / "out"

            headers = [
                "Ignore0",
                "Desk",
                "Agent",
                "Ignore3",
                "CID",
                "Status",
                "Campaign",
                "Ignore7",
                "Country",
                "Ignore9",
                "Ignore10",
                "Ignore11",
                "Ignore12",
                "Assigned",
                "FTD",
            ]
            rows = [
                ["", "TR1-EN", "Agent A", "", "CID-1", "Reached", "Aff A", "", "Switzerland", "", "", "", "", "1", "0"],
                ["", "TR1-EN", "Agent A", "", "CID-2", "Anything", "Aff A", "", "Switzerland", "", "", "", "", "1", "1"],
                ["", "TR1-EN", "Agent A", "", "CID-3", "No Answer", "Aff A", "", "Switzerland", "", "", "", "", "8", "0"],
                ["", "TR1-EN", "Agent A", "", "CID-4", "Other", "Aff B", "", "Switzerland", "", "", "", "", "0", "1"],
            ]
            _write_input_with_header_row_3(input_path, headers, rows)

            outputs = lead_splitter.build_outputs(
                input_path=input_path,
                output_dir=output_dir,
                generate_lead=False,
                generate_aff=False,
                generate_countries=True,
            )

            self.assertIn("countries", outputs)
            countries_output = outputs["countries"]
            workbook = load_workbook(countries_output, data_only=False)
            sheet = workbook["Lead splitter by countries"]

            self.assertEqual(
                [sheet.cell(1, col).value for col in range(1, 8)],
                ["Desk", "Country", "Campaign", "Status", "Leads", "FTD", "CR"],
            )

            def _find_row(label: str, candidate_cols: list[int]) -> tuple[int, int]:
                for col_idx in candidate_cols:
                    for row_idx in range(2, sheet.max_row + 1):
                        if str(sheet.cell(row_idx, col_idx).value or "").strip() == label:
                            return row_idx, col_idx
                self.fail(f"Could not find '{label}' in columns {candidate_cols}.")

            # Status CR uses status leads / campaign leads.
            no_answer_row, no_answer_col = _find_row("No Answer", [4, 12, 20])
            self.assertAlmostEqual(
                float(sheet.cell(no_answer_row, no_answer_col + 3).value or 0.0),
                0.0,
                places=6,
            )

            reached_row, reached_col = _find_row("Reached", [4, 12, 20])
            self.assertAlmostEqual(
                float(sheet.cell(reached_row, reached_col + 3).value or 0.0),
                0.5,
                places=6,
            )

            # FTD row is turned into Telemarketing and still follows status-leads ratio.
            telemarketing_row, telemarketing_col = _find_row("Telemarketing", [4, 12, 20])
            self.assertAlmostEqual(
                float(sheet.cell(telemarketing_row, telemarketing_col + 3).value or 0.0),
                0.5,
                places=6,
            )

            # Total CR rows use strict FTD / Leads.
            campaign_total_row, campaign_total_col = _find_row("Aff A Total", [3, 11, 19])
            self.assertAlmostEqual(
                float(sheet.cell(campaign_total_row, campaign_total_col + 4).value or 0.0),
                0.5,
                places=6,
            )
            campaign_b_total_row, campaign_b_total_col = _find_row("Aff B Total", [3, 11, 19])
            self.assertAlmostEqual(
                float(sheet.cell(campaign_b_total_row, campaign_b_total_col + 4).value or 0.0),
                0.0,
                places=6,
            )
            country_total_row, country_total_col = _find_row("Switzerland Total", [2, 10, 18])
            self.assertAlmostEqual(
                float(sheet.cell(country_total_row, country_total_col + 5).value or 0.0),
                1.0,
                places=6,
            )


class LeadSplitterArVariantTests(unittest.TestCase):
    def test_ar_desk_and_country_helpers(self) -> None:
        self.assertEqual(lead_splitter.get_desk_ar("AR1-PT"), "AR1")
        self.assertEqual(lead_splitter.get_country_ar("AR1-PT"), "PT")
        self.assertEqual(lead_splitter.get_desk_ar("AR"), "AR")
        self.assertEqual(lead_splitter.get_country_ar("AR"), "")

    def test_ar_lane_assignment(self) -> None:
        lane = lead_splitter.AR_VARIANT.lane_of
        self.assertEqual(lane("AR1"), "left")
        self.assertEqual(lane("AR2"), "middle")
        self.assertEqual(lane("AR"), "middle")
        self.assertEqual(lane("TR"), "right")

    def _ar_headers(self) -> list[str]:
        return [
            "Ignore0",
            "Desk",
            "Agent",
            "Ignore3",
            "CID",
            "Status",
            "Campaign",
            "Ignore7",
            "Country",
            "Ignore9",
            "Ignore10",
            "Ignore11",
            "Ignore12",
            "Assigned",
            "FTD",
        ]

    def test_ar_countries_output_uses_desk_and_suffix_country_and_lanes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "lead_input.xlsx"
            output_dir = root / "out"

            rows = [
                # Country column (index 8) is intentionally wrong to prove AR ignores it.
                ["", "AR1-PT", "Agent A TR", "", "CID-1", "Reached", "Camp A", "", "WRONG", "", "", "", "", "1", "0"],
                ["", "AR2-ES", "Agent B TR", "", "CID-2", "Reached", "Camp B", "", "WRONG", "", "", "", "", "1", "1"],
                ["", "TR-DE", "Agent C TR", "", "CID-3", "Reached", "Camp C", "", "WRONG", "", "", "", "", "1", "0"],
            ]
            _write_input_with_header_row_3(input_path, self._ar_headers(), rows)

            outputs = lead_splitter.build_outputs(
                input_path=input_path,
                output_dir=output_dir,
                generate_lead=False,
                generate_aff=False,
                generate_countries=True,
                variant="ar",
            )
            workbook = load_workbook(outputs["countries"], data_only=False)
            sheet = workbook["Lead splitter by countries"]

            def _cell(row, col):
                return str(sheet.cell(row, col).value or "").strip()

            # Left lane (cols 1-7) -> AR1 desk, country PT (from the suffix, not "WRONG").
            self.assertEqual(_cell(2, 1), "AR1")
            self.assertEqual(_cell(2, 2), "PT")
            # Middle lane (cols 9-15) -> AR2 desk.
            self.assertEqual(_cell(2, 9), "AR2")
            self.assertEqual(_cell(2, 10), "ES")
            # Right lane (cols 17-23) -> TR desk.
            self.assertEqual(_cell(2, 17), "TR")
            self.assertEqual(_cell(2, 18), "DE")

    def test_ar_agent_name_strips_trailing_tr(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "lead_input.xlsx"
            output_dir = root / "out"

            rows = [
                ["", "AR1-PT", "Mohammed TR", "", "CID-1", "Reached", "Camp A", "", "PT", "", "", "", "", "1", "0"],
            ]
            _write_input_with_header_row_3(input_path, self._ar_headers(), rows)

            outputs = lead_splitter.build_outputs(
                input_path=input_path,
                output_dir=output_dir,
                generate_lead=True,
                generate_aff=False,
                generate_countries=False,
                variant="ar",
            )
            workbook = load_workbook(outputs["lead"], data_only=False)
            pivot = workbook["Pivot"]
            agents = {
                str(pivot.cell(r, 3).value or "").strip()
                for r in range(2, pivot.max_row + 1)
            }
            self.assertIn("Mohammed", agents)
            self.assertNotIn("Mohammed TR", agents)

    def test_ar_aff_tables_are_by_desk(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "lead_input.xlsx"
            output_dir = root / "out"

            rows = [
                ["", "AR1-PT", "Agent A TR", "", "CID-1", "Reached", "Camp A", "", "PT", "", "", "", "", "1", "1"],
                ["", "AR2-ES", "Agent B TR", "", "CID-2", "Reached", "Camp B", "", "ES", "", "", "", "", "1", "0"],
                ["", "TR-DE", "Agent C TR", "", "CID-3", "Reached", "Camp C", "", "DE", "", "", "", "", "1", "0"],
            ]
            _write_input_with_header_row_3(input_path, self._ar_headers(), rows)

            outputs = lead_splitter.build_outputs(
                input_path=input_path,
                output_dir=output_dir,
                generate_lead=False,
                generate_aff=True,
                generate_countries=False,
                variant="ar",
            )
            workbook = load_workbook(outputs["aff"], data_only=False)
            sheet = workbook["AFF by Status"]

            # Header first column is "Desk" and the three tables are labeled by desk.
            self.assertEqual(str(sheet.cell(1, 1).value or "").strip(), "Desk")
            self.assertEqual(str(sheet.cell(2, 1).value or "").strip(), "AR1")
            self.assertEqual(str(sheet.cell(2, 9).value or "").strip(), "AR2")
            self.assertEqual(str(sheet.cell(2, 17).value or "").strip(), "TR")


if __name__ == "__main__":
    unittest.main()
