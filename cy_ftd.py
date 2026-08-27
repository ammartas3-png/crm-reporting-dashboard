"""Add CY-desk FTD leads to the report output as Telemarketing rows.

The CY FTDs file is a separate Excel export whose headers start on the 3rd row.
Only rows where column O (``FTD Cnt``) equals 1 and column B (``Desk``) starts
with ``CY`` are considered. Each such lead is added to the report output as a
Telemarketing row, unless the same Account No + Brand already exists in the CRM
output as a Telemarketing row. The added rows then go through the same KYC
comment lookup as every other Telemarketing row.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CY_FTD_HEADER_ROW = 3
DATE_OF_BIRTH_OUTPUT_COLUMN = "Date of birth"

# Needed columns -> acceptable header names (normalized: lowercased, single-spaced).
_COLUMN_CANDIDATES: dict[str, list[str]] = {
    "brand": ["brand"],
    "desk": ["desk"],
    "account_no": ["account no", "account no.", "account number"],
    "ftd_cnt": ["ftd cnt", "ftd count"],
    "owner": ["ftd transaction owner"],
    "campaign": ["campaign"],
    "country": ["country"],
    "registration_time": [
        "registration time(utc)",
        "registration time (utc)",
        "registration time",
    ],
}

_COLUMN_LABELS = {
    "brand": "Brand",
    "desk": "Desk",
    "account_no": "Account No",
    "ftd_cnt": "FTD Cnt",
    "owner": "FTD Transaction Owner",
    "campaign": "Campaign",
    "country": "Country",
    "registration_time": "Registration Time(UTC)",
}

_OWNER_CODE_RE = re.compile(r"\s*\([^)]*\)\s*$")


def _normalize_header(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text if text is not None else "").strip()).casefold()


def strip_owner_code(name: Any) -> str:
    """'Marco Mo (PW1020)' -> 'Marco Mo'."""
    return _OWNER_CODE_RE.sub("", str(name if name is not None else "")).strip()


def _is_ftd_one(value: Any) -> bool:
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        try:
            return float(value) == 1.0
        except (TypeError, ValueError):
            return False
    text = str(value).strip()
    if not text:
        return False
    try:
        return float(text) == 1.0
    except ValueError:
        return text == "1"


def _desk_is_cy(desk: Any) -> bool:
    return str(desk if desk is not None else "").strip().upper().startswith("CY")


def _resolve_columns(header_row: tuple[Any, ...]) -> dict[str, int]:
    normalized = {}
    for idx, header in enumerate(header_row):
        key = _normalize_header(header)
        if key and key not in normalized:
            normalized[key] = idx

    resolved: dict[str, int | None] = {}
    for key, candidates in _COLUMN_CANDIDATES.items():
        found: int | None = None
        for candidate in candidates:
            if candidate in normalized:
                found = normalized[candidate]
                break
        if found is None and key == "registration_time":
            for header_key, idx in normalized.items():
                if "registration time" in header_key:
                    found = idx
                    break
        resolved[key] = found

    missing = [_COLUMN_LABELS[key] for key, idx in resolved.items() if idx is None]
    if missing:
        raise ValueError(
            "CY FTDs file is missing required columns: " + ", ".join(missing)
        )
    return {key: idx for key, idx in resolved.items() if idx is not None}


def read_cy_ftd_entries(path: Path) -> list[dict[str, Any]]:
    """Return the qualifying CY-desk FTD leads from the CY FTDs Excel file."""
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if len(rows) < CY_FTD_HEADER_ROW:
        return []

    columns = _resolve_columns(rows[CY_FTD_HEADER_ROW - 1])

    def _cell(row: tuple[Any, ...], key: str) -> Any:
        idx = columns[key]
        return row[idx] if idx < len(row) else None

    entries: list[dict[str, Any]] = []
    for row in rows[CY_FTD_HEADER_ROW:]:
        if row is None or all(value is None for value in row):
            continue
        if not _is_ftd_one(_cell(row, "ftd_cnt")):
            continue
        if not _desk_is_cy(_cell(row, "desk")):
            continue
        entries.append(
            {
                "brand": _cell(row, "brand"),
                "desk": _cell(row, "desk"),
                "owner": strip_owner_code(_cell(row, "owner")),
                "account_no": _cell(row, "account_no"),
                "campaign": _cell(row, "campaign"),
                "country": _cell(row, "country"),
                "registration_time": _cell(row, "registration_time"),
            }
        )
    return entries


def _normalize_match_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(value) if value.is_integer() else value
    text = str(value).strip()
    text = re.sub(r"\.0+$", "", text)
    return text.casefold()


def _normalize_match_platform(value: Any) -> str:
    return str(value if value is not None else "").strip().casefold()


def _match_key(account_no: Any, brand: Any) -> tuple[str, str]:
    return (_normalize_match_id(account_no), _normalize_match_platform(brand))


def existing_telemarketing_keys(rows: list[dict[str, Any]]) -> set[tuple[str, str]]:
    keys: set[tuple[str, str]] = set()
    for row in rows:
        status = str(row.get("Status", "") or "").strip().casefold()
        if status == "telemarketing":
            keys.add(_match_key(row.get("ID"), row.get("Platform")))
    return keys


def _build_row(entry: dict[str, Any]) -> dict[str, Any]:
    return {
        "Platform": entry.get("brand"),
        "Customer Type": "",
        "ID": entry.get("account_no"),
        "Created": entry.get("registration_time"),
        "Name": "",
        "Department": entry.get("desk"),
        "Status": "Telemarketing",
        "CB": None,
        "Country": entry.get("country"),
        "Campaign": entry.get("campaign"),
        "Sub-Campaign": "",
        "Placement": "",
        "Assigned to": entry.get("owner"),
        DATE_OF_BIRTH_OUTPUT_COLUMN: "",
        "Comments": "",
        "_comments_yellow": False,
        "Call Attempts": 1,
    }


def apply_cy_ftd_rows(all_rows: list[dict[str, Any]], cy_ftd_report: Path | None) -> None:
    """Append CY-desk FTD Telemarketing rows to ``all_rows`` (in place).

    Rows whose Account No + Brand already appear as a Telemarketing row are
    skipped so no duplicates are created.
    """
    if cy_ftd_report is None:
        return

    entries = read_cy_ftd_entries(cy_ftd_report)
    if not entries:
        return

    seen = existing_telemarketing_keys(all_rows)
    for entry in entries:
        key = _match_key(entry.get("account_no"), entry.get("brand"))
        if key in seen:
            continue
        seen.add(key)
        all_rows.append(_build_row(entry))
